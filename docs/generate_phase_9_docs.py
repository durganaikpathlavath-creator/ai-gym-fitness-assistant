import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def generate_phase_9_docx():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Segoe UI'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)  # Slate 800

    # Title
    title = doc.add_paragraph()
    r_title = title.add_run("Phase 9 — Dashboards & Advanced Analytics")
    r_title.font.name = 'Segoe UI Semibold'
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = RGBColor(0x6D, 0x28, 0xD9)  # Purple 700

    sub = doc.add_paragraph()
    r_sub = sub.add_run("Master Implementation Guide & Architectural Specification")
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)  # Slate 500

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Callout Box
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F5F3FF")  # Purple 50
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    p_call = cell.paragraphs[0]
    r_call_title = p_call.add_run("ANALYTICS ENGINE POLICY & DATA TRUTHFULNESS\n")
    r_call_title.bold = True
    r_call_title.font.size = Pt(10)
    r_call_title.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)

    r_call_text = p_call.add_run(
        "All analytical metrics displayed across the dashboard are strictly grounded in authenticated "
        "database records. Data fabrication is explicitly prohibited. Time-window parameters (7d/14d/30d) "
        "aggregate rolling historical metrics cleanly. Insufficient data states are returned with explicit "
        "has_data=false flags, and cross-domain insights are explicitly designated as observational non-causal correlations."
    )
    r_call_text.font.size = Pt(9.5)
    r_call_text.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Section 1
    h1 = doc.add_paragraph()
    r_h1 = h1.add_run("1. Executive Summary & Core Requirements")
    r_h1.font.name = 'Segoe UI Semibold'
    r_h1.font.size = Pt(15)
    r_h1.font.color.rgb = RGBColor(0x4C, 0x1D, 0x95)

    doc.add_paragraph(
        "Phase 9 establishes the project's comprehensive Analytics & Dashboard layer. "
        "It integrates data from Phase 1 (Profiles), Phase 2 (Workout Sessions), Phase 3 (Performance Scores), "
        "Phase 4 (Nutrition Logs & Targets), Phase 6 (Habit Skip Risk), and Phase 8 (Smart Gym IoT Telemetry)."
    )

    # Section 2 Table of Metrics
    h2 = doc.add_paragraph()
    r_h2 = h2.add_run("2. Derived Analytics & Formula Specifications")
    r_h2.font.name = 'Segoe UI Semibold'
    r_h2.font.size = Pt(15)
    r_h2.font.color.rgb = RGBColor(0x4C, 0x1D, 0x95)

    table_m = doc.add_table(rows=1, cols=4)
    table_m.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table_m.rows[0].cells
    hdr_titles = ["Metric Name", "Domain Source", "Mathematical Formula", "Target Interpretation"]
    for i, title_text in enumerate(hdr_titles):
        set_cell_background(hdr_cells[i], "6D28D9")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(title_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)

    metrics_data = [
        ("Total Workouts", "Phase 2 Workouts", "COUNT(WorkoutSession.id)", "Total completed sessions in window"),
        ("Average Score", "Phase 3 Performance", "SUM(performance_score) / N", "Mean pose quality score [0-100]"),
        ("Performance Trend", "Phase 3 Performance", "Phase 3 PerformanceService", "improving (>= +2.5) / declining (<= -2.5) / stable"),
        ("Calorie Compliance", "Phase 4 Nutrition", "(Compliant Days / Logged Days) * 100", "Compliant if intake within 85%–110% of target"),
        ("Consistency Rate %", "Phase 6 Habits", "min(100.0, (Actual / Expected) * 100)", "Expected = (Target Days / 7) * Window"),
        ("Avg IoT Intensity", "Phase 8 Smart Gym", "SUM(intensity_score) / N_telem", "Mean set intensity % across devices"),
    ]

    for m_name, m_dom, m_form, m_interp in metrics_data:
        row_cells = table_m.add_row().cells
        for i, val in enumerate([m_name, m_dom, m_form, m_interp]):
            set_cell_background(row_cells[i], "F8FAFC")
            set_cell_margins(row_cells[i], top=80, bottom=80, left=100, right=100)
            p = row_cells[i].paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Save document
    out_path = os.path.join(os.path.dirname(__file__), "PHASE_9_MASTER_IMPLEMENTATION_GUIDE.docx")
    doc.save(out_path)
    print(f"Generated Phase 9 docx at: {out_path}")


if __name__ == "__main__":
    generate_phase_9_docx()

