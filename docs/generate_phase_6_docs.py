import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_callout_box(doc, text, title="BEHAVIORAL AI SAFETY & NON-CAUSAL GUIDELINES", bg_hex="EEF2FF", border_hex="4F46E5"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24')  # 3pt
    left.set(qn('w:space'), '0')
    left.set(qn('w:color'), border_hex)
    tcBorders.append(left)

    for side in ['top', 'bottom', 'right']:
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:val'), 'none')
        tcBorders.append(node)

    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"🧠 {title}\n")
    run_t.bold = True
    run_t.font.size = Pt(11)
    run_t.font.color.rgb = RGBColor(79, 70, 229)

    run_b = p.add_run(text)
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = RGBColor(30, 41, 59)


def generate_docx():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("AI Gym & Fitness Assistant")
    run_title.font.size = Pt(26)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(79, 70, 229)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    run_sub = p_sub.add_run("Phase 6 — Master Implementation Guide: Fitness Habit Tracker & Behavioral AI")
    run_sub.font.size = Pt(15)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Callout Box
    add_callout_box(
        doc,
        "The Phase 6 Fitness Habit Tracker operates strictly as a behavioral feedback system. "
        "All factor attributions use explicit non-causal language ('associated with', 'correlated with') "
        "and predictions are strictly grounded in user-logged workout history to ensure accuracy and prevent hallucination.",
        title="BEHAVIORAL AI DESIGN & NON-CAUSAL SAFETY POLICY"
    )

    doc.add_paragraph()

    # 1. Executive Summary
    h1 = doc.add_heading("1. Executive Summary", level=1)
    h1.runs[0].font.color.rgb = RGBColor(79, 70, 229)
    
    p = doc.add_paragraph(
        "Phase 6 introduces the Fitness Habit Tracker & Behavioral AI module into the AI Gym & Fitness Assistant. "
        "By leveraging actual historical workout completion data, the system extracts an 8-feature behavioral telemetry vector, "
        "trains a scikit-learn Logistic Regression model with temporal time-series split validation, predicts the probability "
        "of a user skipping their upcoming workout window, provides non-causal factor attributions, generates adaptive nudges, "
        "and computes evidence-based schedule recommendations."
    )
    p.paragraph_format.space_after = Pt(12)

    # 2. Key Technical Components
    h2 = doc.add_heading("2. Core Architectural Features", level=1)
    h2.runs[0].font.color.rgb = RGBColor(79, 70, 229)

    components = [
        ("HabitFeatureEngine", "Extracts 8 deterministic behavioral features (days_since_last_workout, 7d/30d frequencies, avg weekly, consistency score, preferred weekday ratio, max gap 30d, form score trend delta) with strict upper cutoff boundaries."),
        ("Strict Temporal Target Cutoff", "Enforces target_end <= evaluation_cutoff for every historical sample in dataset construction, guaranteeing zero future leakage."),
        ("HabitPredictorService", "Trains scikit-learn LogisticRegression (C=1.0, max_iter=1000) using temporal observation windows. Uses deterministic SparseDataFallbackModel when labeled training samples are sparse or lack class variance."),
        ("Cold-Start vs ML Disambiguation", "Enforces strict handling for users with < 3 sessions (status: 'insufficient_data', skip_probability: null), while distinguishing ML-trained models (ml_trained: true) from fallback models (ml_trained: false)."),
        ("Non-Causal Attributions & Observational Schedule", "Generates protective or risk-elevating factor attributions and observational schedule intelligence avoiding absolute optimality claims."),
        ("Habit Tracker UI", "Next.js dashboard page featuring a skip risk probability gauge, adaptive nudge banner, schedule recommendation card, 8-feature telemetry grid, and historical timeline table.")
    ]

    for title, desc in components:
        p = doc.add_paragraph()
        run_b = p.add_run(f"• {title}: ")
        run_b.bold = True
        run_b.font.color.rgb = RGBColor(30, 41, 59)
        run_d = p.add_run(desc)
        run_d.font.color.rgb = RGBColor(71, 85, 105)
        p.paragraph_format.space_after = Pt(4)

    doc.add_paragraph()

    # 3. Database Schema Table
    h3 = doc.add_heading("3. Database Schema (habit_predictions)", level=1)
    h3.runs[0].font.color.rgb = RGBColor(79, 70, 229)

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    headers = ["Column Name", "Type", "Constraints", "Description"]
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "4F46E5")
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)

    schema_data = [
        ("id", "INTEGER", "PRIMARY KEY", "Auto-incrementing record ID"),
        ("user_id", "INTEGER", "FK -> users.id (CASCADE)", "Linked user account ID"),
        ("prediction_date", "TIMESTAMP", "NOT NULL", "Snapshot prediction evaluation date"),
        ("skip_probability", "FLOAT", "NULLABLE", "Predicted skip probability [0.00, 1.00]"),
        ("risk_level", "VARCHAR(30)", "NOT NULL", "low, moderate, high, or insufficient_data"),
        ("primary_factor", "VARCHAR(255)", "NULLABLE", "Primary behavioral signal key"),
        ("nudge_text", "TEXT", "NULLABLE", "Grounded behavioral nudge text"),
        ("recommended_schedule", "VARCHAR(255)", "NULLABLE", "Evidence-based schedule recommendation"),
        ("created_at", "TIMESTAMP", "NOT NULL", "Telemetry creation timestamp"),
    ]

    for row_idx, row_data in enumerate(schema_data):
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = val
            set_cell_background(row_cells[i], "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(row_cells[i], top=100, bottom=100, left=150, right=150)
            p = row_cells[i].paragraphs[0]
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph()

    # 4. Automated Verification Results
    h4 = doc.add_heading("4. Automated Verification & Regression Suite", level=1)
    h4.runs[0].font.color.rgb = RGBColor(79, 70, 229)

    verif_items = [
        ("Phase 6 Test Suite (test_phase_6_habit.py)", "22/22 PASSED (0.710s)"),
        ("Phase 1 User Auth Suite (test_phase_1_users.py)", "9/9 PASSED (1.486s)"),
        ("Phase 2 Vision API Suite (test_phase_2_api.py)", "4/4 PASSED (0.607s)"),
        ("Phase 3 Weekly Performance Suite (test_phase_3_performance.py)", "10/10 PASSED (0.571s)"),
        ("Phase 4 Nutrition Suite (test_phase_4_nutrition.py)", "13/13 PASSED (0.759s)"),
        ("Phase 5 Gym Buddy Suite (test_phase_5_buddy.py)", "13/13 PASSED (1.299s)"),
        ("Next.js Production Build (npm run build)", "0 ERRORS (13 static pages compiled)")
    ]

    for test_name, status in verif_items:
        p = doc.add_paragraph()
        run_t = p.add_run(f"✅ {test_name}: ")
        run_t.bold = True
        run_t.font.color.rgb = RGBColor(16, 185, 129)
        run_s = p.add_run(status)
        run_s.font.color.rgb = RGBColor(30, 41, 59)
        p.paragraph_format.space_after = Pt(4)

    # Save document
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "PHASE_6_MASTER_IMPLEMENTATION_GUIDE.docx")
    doc.save(output_path)
    print(f"Generated Word document successfully: {output_path}")


if __name__ == "__main__":
    generate_docx()

