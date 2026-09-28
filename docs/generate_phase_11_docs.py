import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def generate_phase_11_docx():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Segoe UI'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)  # Slate 800

    # Title
    title = doc.add_paragraph()
    r_title = title.add_run("Phase 11 — Full System Integration + Security + Testing")
    r_title.font.name = 'Segoe UI Semibold'
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)  # Indigo 600

    sub = doc.add_paragraph()
    r_sub = sub.add_run("Master Implementation Guide & Verification Report")
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)  # Slate 500

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Callout Box
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "EEF2FF")  # Indigo 50
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    p_call = cell.paragraphs[0]
    r_call_title = p_call.add_run("SYSTEM INTEGRATION & SECURITY ASSURANCE\n")
    r_call_title.bold = True
    r_call_title.font.size = Pt(10)
    r_call_title.font.color.rgb = RGBColor(0x37, 0x30, 0xA3)  # Indigo 800

    r_call_text = p_call.add_run(
        "Phase 11 unifies the independently built Phases 1 through 10 into an integrated fitness application. "
        "Comprehensive end-to-end integration workflows, cross-user data isolation tests across all 10 domain modules, "
        "JWT identity scoping audits, SQL injection resilience, error information leakage prevention, database ON DELETE CASCADE "
        "integrity checks, and Next.js static production build validation have been fully executed and verified."
    )
    r_call_text.font.size = Pt(9.5)
    r_call_text.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Reading markdown content
    md_path = os.path.join(os.path.dirname(__file__), "PHASE_11_MASTER_IMPLEMENTATION_GUIDE.md")
    if os.path.exists(md_path):
        with open(md_path, "r", encoding="utf-8") as f:
            lines = f.readlines()

        in_code_block = False
        code_lines = []

        for line in lines:
            stripped = line.strip()

            if stripped.startswith("```"):
                if in_code_block:
                    code_text = "".join(code_lines)
                    c_table = doc.add_table(rows=1, cols=1)
                    c_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    c_cell = c_table.cell(0, 0)
                    set_cell_background(c_cell, "F8FAFC")  # Slate 50
                    set_cell_margins(c_cell, top=100, bottom=100, left=150, right=150)
                    p_code = c_cell.paragraphs[0]
                    r_code = p_code.add_run(code_text.rstrip())
                    r_code.font.name = 'Consolas'
                    r_code.font.size = Pt(9)
                    r_code.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                    doc.add_paragraph().paragraph_format.space_after = Pt(8)
                    code_lines = []
                    in_code_block = False
                else:
                    in_code_block = True
                    code_lines = []
                continue

            if in_code_block:
                code_lines.append(line)
                continue

            if stripped.startswith("# "):
                h = doc.add_paragraph()
                r = h.add_run(stripped[2:])
                r.font.name = 'Segoe UI Semibold'
                r.font.size = Pt(18)
                r.font.color.rgb = RGBColor(0x37, 0x30, 0xA3)
                h.paragraph_format.space_before = Pt(16)
                h.paragraph_format.space_after = Pt(6)
            elif stripped.startswith("## "):
                h = doc.add_paragraph()
                r = h.add_run(stripped[3:])
                r.font.name = 'Segoe UI Semibold'
                r.font.size = Pt(14)
                r.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
                h.paragraph_format.space_before = Pt(12)
                h.paragraph_format.space_after = Pt(4)
            elif stripped.startswith("### "):
                h = doc.add_paragraph()
                r = h.add_run(stripped[4:])
                r.font.name = 'Segoe UI Semibold'
                r.font.size = Pt(12)
                r.font.color.rgb = RGBColor(0x43, 0x38, 0xCA)
                h.paragraph_format.space_before = Pt(8)
                h.paragraph_format.space_after = Pt(2)
            elif stripped.startswith("- "):
                p = doc.add_paragraph(style='List Bullet')
                r = p.add_run(stripped[2:])
                r.font.name = 'Segoe UI'
                r.font.size = Pt(10)
            elif stripped.startswith("|"):
                p = doc.add_paragraph()
                r = p.add_run(stripped)
                r.font.name = 'Consolas'
                r.font.size = Pt(9)
            elif stripped:
                p = doc.add_paragraph()
                r = p.add_run(stripped)
                r.font.name = 'Segoe UI'
                r.font.size = Pt(10.5)
                p.paragraph_format.space_after = Pt(4)

    out_docx = os.path.join(os.path.dirname(__file__), "PHASE_11_MASTER_IMPLEMENTATION_GUIDE.docx")
    doc.save(out_docx)
    print(f"Successfully generated Phase 11 DOCX documentation: {out_docx}")


if __name__ == "__main__":
    generate_phase_11_docx()

